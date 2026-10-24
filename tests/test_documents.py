"""文档上传与向量化管线测试。

- 解析器：TXT/DOCX 正常解析、损坏 PDF 报错；
- 接口：类型/大小校验、异步完成状态流转、分块入库与向量检索、删除级联、越权隔离、
  未配置 API Key 时失败并回写原因且不留脏分块。

真实库接口测试需要配置 ``DATABASE_URL``（与 test_auth.py 一致），未配置时自动跳过。
"""

import uuid
from dataclasses import replace
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document as DocxDocument
from httpx import ASGITransport, AsyncClient
from sqlalchemy import delete, func, select

from app.config import settings
from app.database import async_session_factory, engine, init_database
from app.main import app
from app.models import Document, DocumentChunk, User
from app.services import auth_service, document_service, embedding_client
from app.services.document_parser import DocumentParseError, parse_document_text

requires_db = pytest.mark.skipif(
    not settings.database_url,
    reason="未配置 DATABASE_URL，跳过真实库测试",
)


# ---------- 解析器 ----------


def test_parse_docx_paragraphs(tmp_path: Path) -> None:
    """DOCX 段落文本可正常提取。"""
    path = tmp_path / "sample.docx"
    document = DocxDocument()
    document.add_paragraph("第一段内容")
    document.add_paragraph("")
    document.add_paragraph("第二段内容")
    document.save(path)

    text = parse_document_text(path, "docx")

    assert "第一段内容" in text
    assert "第二段内容" in text


def test_parse_txt_gb18030(tmp_path: Path) -> None:
    """TXT 按 GB18030 编码时可正常解码。"""
    path = tmp_path / "sample.txt"
    path.write_bytes("中文内容测试".encode("gb18030"))

    assert parse_document_text(path, "txt") == "中文内容测试"


def test_parse_corrupt_pdf_raises(tmp_path: Path) -> None:
    """损坏的 PDF 抛 DocumentParseError，不产生静默空结果。"""
    path = tmp_path / "broken.pdf"
    path.write_bytes(b"not a real pdf")

    with pytest.raises(DocumentParseError):
        parse_document_text(path, "pdf")


def test_parse_empty_text_raises(tmp_path: Path) -> None:
    """空白文本文件视为无有效文本。"""
    path = tmp_path / "empty.txt"
    path.write_bytes(b"   \n\n")

    with pytest.raises(DocumentParseError):
        parse_document_text(path, "txt")


# ---------- 真实库接口 ----------


@pytest.fixture
def jwt_settings(monkeypatch: pytest.MonkeyPatch):
    """为 auth_service 注入测试用 JWT 配置。"""
    test_settings = replace(settings, jwt_secret="test-jwt-secret", jwt_expire_minutes=60)
    monkeypatch.setattr(auth_service, "settings", test_settings)
    return test_settings


@pytest.fixture
def upload_settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    """上传目录指向临时目录，大小上限降为 1MB 便于覆盖超限路径。"""
    test_settings = replace(
        settings,
        upload_dir=str(tmp_path / "uploads"),
        upload_max_mb=1,
    )
    monkeypatch.setattr(document_service, "settings", test_settings)
    return Path(test_settings.upload_dir)


@pytest.fixture
def fake_embeddings(monkeypatch: pytest.MonkeyPatch):
    """Mock Embedding：每个分块返回单位向量，便于按余弦距离检索验证。"""

    async def fake_embed_texts(texts: list[str]) -> list[list[float]]:
        vectors: list[list[float]] = []
        for index, _ in enumerate(texts):
            vector = [0.0] * settings.embedding_dimension
            vector[index % settings.embedding_dimension] = 1.0
            vectors.append(vector)
        return vectors

    monkeypatch.setattr(embedding_client, "embed_texts", fake_embed_texts)


@pytest.fixture
async def client():
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as async_client:
        yield async_client


@pytest.fixture
async def test_user(jwt_settings):
    """创建测试用户并签发 token，测试后按用户级联清理数据。"""
    await init_database()
    assert async_session_factory is not None

    username = f"doc_test_{uuid.uuid4().hex[:12]}"
    async with async_session_factory() as session:
        user = User(username=username, password_hash="test-hash")
        session.add(user)
        await session.commit()
        await session.refresh(user)
        user_id = user.id

    token, _ = auth_service.create_access_token(
        User(id=user_id, username=username, password_hash="test-hash")
    )
    yield {"user_id": user_id, "headers": {"Authorization": f"Bearer {token}"}}

    async with async_session_factory() as session:
        await session.execute(delete(User).where(User.id == user_id))
        await session.commit()

    # 每个 async 测试运行在独立事件循环，释放连接池避免跨循环复用残留连接
    if engine is not None:
        await engine.dispose()


async def _create_kb(client: AsyncClient, headers: dict, name: str) -> int:
    """通过接口创建知识库并返回 kb_id。"""
    response = await client.post("/api/knowledge-bases", json={"name": name}, headers=headers)
    assert response.status_code == 200
    return response.json()["data"]["id"]


async def _upload(
    client: AsyncClient,
    kb_id: int,
    headers: dict,
    filename: str,
    content: bytes,
):
    """上传文件（multipart）并返回响应对象。"""
    return await client.post(
        f"/api/knowledge-bases/{kb_id}/documents",
        files={"file": (filename, content, "application/octet-stream")},
        headers=headers,
    )


def _build_pdf(text: str) -> bytes:
    """构造带文本内容的单页最小 PDF（Helvetica），用于真实 PDF 解析验证。"""
    content = f"BT /F1 18 Tf 72 720 Td ({text}) Tj ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>"
        ),
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream".encode(),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(output))
        output += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"

    xref_offset = len(output)
    output += f"xref\n0 {len(objects) + 1}\n".encode()
    output += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        output += f"{offset:010d} 00000 n \n".encode()
    output += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
        f"startxref\n{xref_offset}\n%%EOF\n"
    ).encode()
    return bytes(output)


@requires_db
async def test_upload_markdown_completes_and_vector_search_works(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
    fake_embeddings,
) -> None:
    """上传 MD 后状态由处理中转完成，分块入库且可按余弦距离检索回最近分块。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "上传测试库")
    content = "# 标题\n\n第一段知识内容。\n\n第二段知识内容。".encode("utf-8")

    response = await _upload(client, kb_id, headers, "笔记.md", content)
    assert response.status_code == 200
    body = response.json()
    assert body["code"] == 0
    assert body["data"]["status"] == "processing"
    assert body["data"]["filename"] == "笔记.md"
    document_id = body["data"]["id"]

    # 后台任务随响应执行完毕：列表中状态为已完成，分块数可见
    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    items = {item["id"]: item for item in list_response.json()["data"]}
    item = items[document_id]
    assert item["status"] == "completed"
    assert item["chunk_count"] >= 1
    assert item["error_message"] is None
    assert len(item["create_time"]) == 19 and item["create_time"][10] == " "

    # 分块文本可查看，序号从 0 递增
    chunks_response = await client.get(f"/api/documents/{document_id}/chunks", headers=headers)
    chunks = chunks_response.json()["data"]
    assert len(chunks) == item["chunk_count"]
    assert [chunk["chunk_index"] for chunk in chunks] == list(range(len(chunks)))
    assert "知识内容" in "".join(chunk["content"] for chunk in chunks)

    # 向量检索：用第一块的向量查询，应命中同一文档 chunk_index=0
    assert async_session_factory is not None
    query_vector = [0.0] * settings.embedding_dimension
    query_vector[0] = 1.0
    async with async_session_factory() as session:
        stmt = (
            select(DocumentChunk)
            .where(DocumentChunk.kb_id == kb_id)
            .order_by(DocumentChunk.embedding.cosine_distance(query_vector))
            .limit(1)
        )
        closest = (await session.execute(stmt)).scalar_one()
        assert closest.document_id == document_id
        assert closest.chunk_index == 0


@requires_db
async def test_upload_real_pdf_parsed_and_chunked(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
    fake_embeddings,
) -> None:
    """上传真实 PDF 后可解析出文本并分块入库。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "PDF 测试库")

    response = await _upload(
        client,
        kb_id,
        headers,
        "sample.pdf",
        _build_pdf("Hello Knowledge Base PDF"),
    )
    assert response.status_code == 200
    document_id = response.json()["data"]["id"]

    chunks_response = await client.get(f"/api/documents/{document_id}/chunks", headers=headers)
    chunks = chunks_response.json()["data"]
    assert chunks
    assert "Hello Knowledge Base PDF" in chunks[0]["content"]


@requires_db
async def test_upload_unsupported_type_rejected(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
) -> None:
    """非法文件类型返回 400 统一错误体，不产生文档记录与落盘文件。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "类型校验库")

    response = await _upload(client, kb_id, headers, "virus.exe", b"not allowed")
    assert response.status_code == 400
    assert response.json()["code"] == 40001

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    assert list_response.json()["data"] == []
    assert not upload_settings.exists() or list(upload_settings.iterdir()) == []


@requires_db
async def test_upload_empty_file_rejected(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
) -> None:
    """空文件直接拒绝，不创建处理中的文档。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "空文件库")

    response = await _upload(client, kb_id, headers, "empty.txt", b"")
    assert response.status_code == 400
    assert response.json()["code"] == 40001

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    assert list_response.json()["data"] == []


@requires_db
async def test_upload_oversize_rejected(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
) -> None:
    """超过大小上限的文件被拒绝（测试中上限降为 1MB）。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "大小校验库")

    oversize = b"a" * (1024 * 1024 + 1)
    response = await _upload(client, kb_id, headers, "big.txt", oversize)
    assert response.status_code == 400
    assert "超过限制" in response.json()["message"]

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    assert list_response.json()["data"] == []


@requires_db
async def test_corrupt_pdf_marks_failed_without_dirty_chunks(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
) -> None:
    """损坏 PDF 解析失败时文档置失败并记录原因，不产生分块数据。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "损坏文件库")

    response = await _upload(client, kb_id, headers, "broken.pdf", b"definitely not a pdf")
    assert response.status_code == 200
    document_id = response.json()["data"]["id"]

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    item = {item["id"]: item for item in list_response.json()["data"]}[document_id]
    assert item["status"] == "failed"
    assert item["error_message"]
    assert item["chunk_count"] == 0

    assert async_session_factory is not None
    async with async_session_factory() as session:
        chunk_count = (
            await session.execute(
                select(func.count(DocumentChunk.id)).where(
                    DocumentChunk.document_id == document_id
                )
            )
        ).scalar_one()
        assert chunk_count == 0


@requires_db
async def test_missing_api_key_marks_failed_without_dirty_chunks(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """未配置 DASHSCOPE_API_KEY 时文档置失败并提示原因，不留下分块数据。"""
    monkeypatch.setattr(
        embedding_client,
        "settings",
        replace(settings, dashscope_api_key=""),
    )
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "未配置 AI 库")

    response = await _upload(client, kb_id, headers, "note.txt", "一些文本内容".encode("utf-8"))
    document_id = response.json()["data"]["id"]

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    item = {item["id"]: item for item in list_response.json()["data"]}[document_id]
    assert item["status"] == "failed"
    assert "AI 服务未配置" in item["error_message"]
    assert item["chunk_count"] == 0

    assert async_session_factory is not None
    async with async_session_factory() as session:
        chunk_count = (
            await session.execute(
                select(func.count(DocumentChunk.id)).where(
                    DocumentChunk.document_id == document_id
                )
            )
        ).scalar_one()
        assert chunk_count == 0


@requires_db
async def test_delete_document_cascades_chunks_and_removes_file(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
    fake_embeddings,
) -> None:
    """删除文档后分块记录清空、磁盘文件删除，列表不再展示。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "删除文档库")

    response = await _upload(client, kb_id, headers, "to_delete.txt", b"delete me please")
    document_id = response.json()["data"]["id"]

    assert async_session_factory is not None
    async with async_session_factory() as session:
        document = await session.get(Document, document_id)
        assert document is not None
        file_path = Path(document.file_path)
    assert file_path.is_file()

    delete_response = await client.delete(f"/api/documents/{document_id}", headers=headers)
    assert delete_response.status_code == 200
    assert delete_response.json()["code"] == 0

    async with async_session_factory() as session:
        chunk_count = (
            await session.execute(
                select(func.count(DocumentChunk.id)).where(
                    DocumentChunk.document_id == document_id
                )
            )
        ).scalar_one()
        assert chunk_count == 0
        assert await session.get(Document, document_id) is None

    assert not file_path.exists()

    list_response = await client.get(f"/api/knowledge-bases/{kb_id}/documents", headers=headers)
    assert list_response.json()["data"] == []


@requires_db
async def test_cannot_access_other_users_document(
    client: AsyncClient,
    test_user: dict,
    upload_settings,
    fake_embeddings,
) -> None:
    """其他用户查看分块/删除文档返回 404 统一错误体。"""
    headers = test_user["headers"]
    kb_id = await _create_kb(client, headers, "越权测试库")
    response = await _upload(client, kb_id, headers, "private.txt", b"private content")
    document_id = response.json()["data"]["id"]

    assert async_session_factory is not None
    other_username = f"doc_test_{uuid.uuid4().hex[:12]}"
    async with async_session_factory() as session:
        other = User(username=other_username, password_hash="test-hash")
        session.add(other)
        await session.commit()
        await session.refresh(other)
        other_id = other.id

    other_token, _ = auth_service.create_access_token(
        User(id=other_id, username=other_username, password_hash="test-hash")
    )
    other_headers = {"Authorization": f"Bearer {other_token}"}

    try:
        chunks_response = await client.get(
            f"/api/documents/{document_id}/chunks",
            headers=other_headers,
        )
        assert chunks_response.status_code == 404
        assert chunks_response.json()["code"] == 40400

        delete_response = await client.delete(
            f"/api/documents/{document_id}",
            headers=other_headers,
        )
        assert delete_response.status_code == 404
        assert delete_response.json()["code"] == 40400
    finally:
        async with async_session_factory() as session:
            await session.execute(delete(User).where(User.id == other_id))
            await session.commit()
